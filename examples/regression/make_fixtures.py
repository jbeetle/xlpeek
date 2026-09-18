# Copyright (c) 2026 henryyu@163.com. All rights reserved.
#
# This file is part of xlpeek, a read-only spreadsheet reader for agents.
# See README.md for what it does and docs/AGENTS.md for how it is meant to be
# driven.

"""Rebuild the fixture suite reads.

`testdata/round2.xlsx` is committed, so nothing here has to run for a suite to
pass. It exists because the committed file has to be reproducible: a fixture
whose provenance is "someone's machine" cannot be re-derived when the shape it
encodes is questioned, and the shapes below are the ones an external review
found the tool answering wrongly. Regenerating it and re-running the suites
proves the file still means what the expectations say it means.

Needs openpyxl — the only step of the regression suite that does. The suites
themselves are stdlib-only.

    python examples/regression/make_fixtures.py
"""
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - a readable failure beats a traceback
    sys.exit("make_fixtures.py needs openpyxl: pip install openpyxl")

REPO = Path(__file__).resolve().parents[2]
TARGET = REPO / "testdata" / "round2.xlsx"
# The second fixture is a shape rather than a dataset: four sheets, each holding
# one habit of a real office workbook that used to change an answer without
# saying so. It is written by the same script because both files have to be
# re-derivable when the shape they encode is questioned.
OFFICE = REPO / "testdata" / "office.xlsx"


def sheet_filters(wb):
    """筛选: three numeric rows, one of them a percentage column."""
    ws = wb.create_sheet("筛选")
    ws.append(["名称", "金额", "比率", "状态"])
    for name, amount, ratio, state in (
        ("A", 100, 0.05, "完成"),
        ("B", 1000, 0.2, "进行"),
        ("C", 2000, 0.2567, "完成"),
    ):
        ws.append([name, amount, ratio, state])
        ws.cell(row=ws.max_row, column=3).number_format = "0.00%"


def sheet_dates(wb):
    """日期: the same day written three ways, plus a number and a text date.

    A date is a number wearing a number format, so the same value reads as
    "2026-01-01 0:00:00", "01-01-26" or "46023" depending on the file. The
    text column is here to hold --dates to its promise: it rewrites dates, and
    a string that merely looks like one is not a date.
    """
    import datetime

    ws = wb.create_sheet("日期")
    ws.append(["日期", "短日期", "时间", "金额", "文本日期"])
    rows = (
        (datetime.datetime(2026, 1, 1), datetime.datetime(2026, 1, 1),
         datetime.datetime(2026, 1, 1, 9, 30), 1234.5, "2026-01-01"),
        (datetime.datetime(2026, 1, 15), datetime.datetime(2026, 1, 15),
         datetime.datetime(2026, 1, 15, 18, 5, 2), 999.99, "2026-01-15"),
    )
    for date, short, stamp, amount, text in rows:
        ws.append([date, short, stamp, amount, text])
        row = ws.max_row
        ws.cell(row=row, column=2).number_format = "mm-dd-yy"
        ws.cell(row=row, column=3).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=row, column=4).number_format = "¥#,##0.00"


def sheet_trailing(wb):
    """尾部空行: two data rows, then formatted rows holding nothing.

    A spreadsheet's last row is its last formatted row, not its last filled
    one, so a border or a fill dragged down past the data leaves the file
    claiming two hundred rows. Paging used to treat those as pending data and
    hand back blank pages with complete:false, which is the failure this sheet
    is here to catch.
    """
    ws = wb.create_sheet("尾部空行")
    ws.append(["名称", "金额"])
    ws.append(["A", 100])
    ws.append(["B", 200])
    for row in range(5, 201):
        ws.cell(row=row, column=1).font = Font()


def sheet_bracketed(wb):
    """括号列名: a parenthesis in the column name, as real reports write it."""
    ws = wb.create_sheet("括号列名")
    ws.append(["金额(万元)", "成本(万元)"])
    ws.append([500, 300])
    ws.append([400, 200])


def sheet_formula(wb):
    """公式: the workbook a --col or an --agg formula is measured against.

    Sized by hand so that every expectation in regress_round3.py can be checked
    on paper, and shaped like the review that asked for the feature: 部门 A holds
    an outlier, so its mean is 2092 while its median is 120, and the two are not
    interchangeable. The rest of the sheet is the awkward cases — a divisor of
    zero, a code the lookup sheet does not hold, and a padded text column.
    """
    ws = wb.create_sheet("公式")
    ws.append(["部门", "金额", "区域", "编号", "分母", "代码"])
    for row in (
        ("A", 100, "华东", " A-1 ", 2, "K1"),
        ("A", 110, "华东", "A-2", 5, "K2"),
        ("A", 120, "华东", "A-3 ", 4, "K3"),
        ("A", 130, "华南", "A-4", 0, "K9"),
        ("A", 10000, "华东", "A-5", 8, "K1"),
        ("B", 200, "华北", " B-1 ", 2, "K2"),
        ("B", 400, "华北", "B-2", 4, "K3"),
        ("B", 600, "华北", "B-3", 5, "K1"),
    ):
        ws.append(list(row))


def sheet_lookup(wb):
    """字典: the table a cross-sheet VLOOKUP reads.

    K9 is deliberately absent, so a lookup of every row has one that finds
    nothing — the case where an empty cell and a failed lookup must not read
    the same.
    """
    ws = wb.create_sheet("字典")
    ws.append(["代码", "目标值"])
    for code, target in (("K1", 11), ("K2", 22), ("K3", 33)):
        ws.append([code, target])


def sheet_high_cardinality(wb):
    """高基数: past the point where grouping stops summarising.

    Grouping by a near-unique column returns the table with extra steps. The
    count is over 1000 on purpose: that is where the default group cap bites.
    """
    ws = wb.create_sheet("高基数")
    ws.append(["凭证号", "金额"])
    for i in range(1200):
        ws.append(["V%04d" % i, i])


def office_hidden_rows(wb):
    """隐藏行: a sheet saved from a filtered or collapsed view.

    The rows a filter removed are still in the file — hidden, which is how Excel
    writes a filtered row. The total on screen is 1600; the total over the file
    is 2100, and until 1.2.0 nothing said which one was being answered.

    C is written with an outline level as well as hidden, which is what a
    collapsed group looks like — the file does not say which of the three ways
    put a row there, and neither does the warning.
    """
    ws = wb.create_sheet("隐藏行")
    ws.append(["项目", "金额"])
    for name, amount in (("A", 100), ("B", 200), ("C", 300),
                         ("D", 400), ("E", 500), ("F", 600)):
        ws.append([name, amount])
    ws.row_dimensions[3].hidden = True            # B, hidden by hand
    ws.row_dimensions[4].hidden = True            # C, collapsed under an outline
    ws.row_dimensions[4].outlineLevel = 1
    ws.auto_filter.ref = "A1:B7"                  # the filter this view was saved from


def office_subtotals(wb):
    """小计行: a report that writes its own totals into the detail column.

    差旅费 100 + 办公费 200 = 小计 300, + 会议费 400 = 合计 700. Summing every
    row gives 1700 — the subtotals counted twice and then counted again — while
    the answer the report itself states is 700.
    """
    ws = wb.create_sheet("小计行")
    ws.append(["科目", "金额"])
    for subject, amount in (("差旅费", 100), ("办公费", 200), ("小计", 300),
                            ("会议费", 400), ("合计", 700)):
        ws.append([subject, amount])


def office_two_level_header(wb):
    """两级表头: a title row merged over the columns it groups.

    The title sits only in the first cell of the merge, which is how Excel
    stores one — so the first column under it is named 2024年金额 by the join,
    and the second needs --fill-merged before it can be named at all. Both
    facts are asserted, because both are what a caller will meet.
    """
    ws = wb.create_sheet("两级表头")
    ws.append(["", "2024年", ""])
    ws.append(["部门", "金额", "数量"])
    for department, amount, count in (("甲", 100, 1), ("乙", 200, 2)):
        ws.append([department, amount, count])
    ws.merge_cells("B1:C1")


def office_full_width(wb):
    """全角数字: what a figure looks like after Word or WeChat.

    Every row is a number to the person reading the sheet; two of the three used
    to be text to every parser, so the column totalled 500 instead of 2968.
    """
    ws = wb.create_sheet("全角数字")
    ws.append(["项目", "金额"])
    for name, amount in (("甲", "１２３４"), ("乙", "１，２３４"), ("丙", 500)):
        ws.append([name, amount])


def build_office():
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in (office_hidden_rows, office_subtotals,
                  office_two_level_header, office_full_width):
        sheet(wb)
    OFFICE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(OFFICE))
    print("wrote %s (%d bytes)" % (OFFICE.relative_to(REPO), OFFICE.stat().st_size))


def build():
    wb = Workbook()
    wb.remove(wb.active)  # the default sheet has no part in the fixture
    for sheet in (sheet_filters, sheet_dates, sheet_trailing,
                  sheet_bracketed, sheet_high_cardinality,
                  sheet_formula, sheet_lookup):
        sheet(wb)
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(TARGET))
    print("wrote %s (%d bytes)" % (TARGET.relative_to(REPO), TARGET.stat().st_size))
    build_office()


if __name__ == "__main__":
    build()
